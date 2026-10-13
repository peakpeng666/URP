"""Post-result narrative and editable Chinese report; no selection tuning."""
import pathlib,json,csv,statistics,hashlib
from docx import Document
from docx.shared import Pt,Inches,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
H=pathlib.Path(__file__).resolve().parents[1]
s=json.loads((H/'analysis/SUMMARY.json').read_bytes());m=s['confirmation'];c=json.loads((H/'CHECK.json').read_bytes());a=json.loads((H/'ANALYSIS_COMPLETE.json').read_bytes());r=json.loads((H/'RUN_COMPLETE.json').read_bytes());mf=json.loads((H/'MODEL_FREEZE.json').read_bytes());assert c['passed']
labels={'weibull_3_45':'Weibull','uniform_1_100':'均匀整数','small_beta_2_5':'偏小Beta','large_beta_5_2':'偏大Beta','normal_50_15':'裁剪正态','bimodal_20_30_70_80':'双峰'}
orders={'random':'随机','ascending':'升序','descending':'降序'}
pnames={'BF':'全程BF','BF_then_EoH':'BF前缀后无条件EoH','published_EoH':'全程固定EoH','SELECTOR':'前缀选择器','SBS':'训练SBS','branch_oracle':'同前缀分支oracle','full_run_oracle':'从头运行的BF/EoH oracle'}
rows=list(csv.DictReader((H/'analysis/family_order_summary.csv').open()));cpurows=list(csv.DictReader((H/'analysis/cpu_accounting.csv').open()));totalcpu=sum(json.loads((H/(stage+'_PROCESS_RESOURCE.json')).read_bytes())['cpu_total_seconds'] for stage in ['RUN','CHECK_CACHED','ANALYZE']);ci=m['cluster_bootstrap_95_percentile_CI'];delta=m['selector_minus_SBS_mean'];head=s['training']['SBS_oracle_headroom'];recovered=m['headroom_recovered_fraction'];full_delta=m['mean_bins']['SELECTOR']-m['mean_bins']['published_EoH']
sections=[]
def sec(title,text):sections.append((title,text))
conclusion='可执行选择收益在本轮预冻的合成混合中成立。' if delta<0 else ('本轮没有把互补性转为优于训练SBS的平均选择收益。' if delta>0 else '本轮选择器与训练SBS平均打平，未产生正的选择收益。')
sec('核心结论',conclusion+f" 确认集108条完整流上，前缀选择器相对训练SBS平均差{delta:+.3f}箱；负数表示选择器节省箱数。训练SBS为{pnames[s['training']['sbs']]}。按同一多重集的三个顺序聚合并在六族内重抽样，95%百分位区间为[{ci[0]:+.3f}, {ci[1]:+.3f}]箱。该区间只描述预设合成混合的不确定性。\n"+f"相对由同一BF前缀后两个可执行续行分支组成的事后oracle，选择器仍有平均{m['selector_oracle_regret']:.3f}箱后悔；SBS的相应余地为{m['SBS_oracle_headroom']:.3f}箱。"+(f" 本选择器收回这部分余地的{100*recovered:.2f}%。" if recovered is not None else ' 余地为零，回收比例不定义。')+f" 逐流胜/平/负={m['W_T_L_vs_SBS'][0]}/{m['W_T_L_vs_SBS'][1]}/{m['W_T_L_vs_SBS'][2]}，最坏相对SBS多{m['worst_harm_vs_SBS']}箱。\n"+f"与从头运行的完整固定EoH相比，选择器平均差{full_delta:+.3f}箱。主指标比较的是共享BF前缀的合法选择任务；完整EoH在前500项也运行EoH，其500项后的箱状态可能不同，仍作为必要固定基线原样列出。")
sec('论文机制与本轮新问题','P069 AutoFolio提供固定库、训练特征预处理、kNN邻域性能选择、训练SBS和最终测试隔离等机制。这里仅落实一个简单选择原型，没有运行其SMAC外层配置，不称整篇复现。\n先前固定EoH压力测试揭示分布和到达顺序的互补，但完整流标签不能用于online选择。本轮只允许观察已到达的前500项及由它们形成的BF残量状态；不允许读取后缀、生成seed、真实分布标签或完整流统计。新训练与新确认种子块均在执行前冻结。')
sec('有限冻结设计','每条流容量100、5000项。算法先用BF处理500项，随后仅决定剩余4500项继续BF还是改用固定发表EoH，不回滚、不重复装箱。所有策略使用5000个初始空槽、全部可行箱、最小索引first argmax；EoH可以在还有可行已用箱时选新箱。\n六个尺寸族为Weibull(形状3,尺度45)、1至100均匀整数、100×Beta(2,5)、100×Beta(5,2)、Normal(50,15)、20至30与70至80的等概率整数双峰。连续值先裁剪到[1,100]再ties-to-even取整。各族训练8个、确认6个全新多重集，每个生成随机、升序、降序三种流。同多重集所有变体同组同split。\n训练144条流，各完整执行BF与BF→EoH，得到288个真实标签。标签必须共享完全一致的BF前缀轨迹，不能拿从头EoH的成绩替代。确认108条流，每条执行全BF、BF→EoH、全EoH和实际选择器，共432次；另有12条风险流的48次。总计768次完整5000步执行。')
sec('选择器与可访问的信息','固定31维特征：10格前缀尺寸频率；前缀均值、总体标准差和10/50/90分位；相邻严格增大与减小比例；后100项与前100项均值差；BF使用箱数/500；10格已用箱残量频率；已用残量均值与总体标准差。尺寸统计均除100，频率用合法前缀计数。\n只用训练集均值/总体标准差作z-score；零标准差置1。固定k=5，平方欧氏距离、等权邻居分别预测两个分支最终箱数，选择预测损失较小者。距离并列采用固定训练行顺序；预测并列退回训练SBS。没有看确认结果后搜索特征、k值或切换点。\n选择模块不导入文件、流对象、随机数或元数据。宿主通过一次性迭代器收取物品，在恰好消费500项、尚未请求第501项时调用选择器，只传已到达前缀、现时残量和已冻结模型。确认选择器完整轨迹还必须与实际选定分支的5000步轨迹一致。')
sec('训练余地与全单元报告',f"训练两个分支的平均箱数为{s['training']['mean_bins'][0]:.3f}与{s['training']['mean_bins'][1]:.3f}；训练分支oracle为{s['training']['branch_oracle_mean']:.3f}，训练SBS的余地为{head:.3f}箱。该余地是额外标签评估后才知道的，已计入训练成本。\n下表完整列出六族×三顺序18个确认单元，每格6个新多重集。相同多重集的不同顺序不是独立抽样；所有18格等权进入预冻主指标。升降序是给定的结构化到达流，不是在线算法获得排序权。没有按结果过滤难例或只叙述胜例。")
sec('同前缀不同未来的风险','四种随机前缀各500项固定不变，后缀均由1至100各45件组成，仅改变4500项后缀的随机、升序、降序。每组3流的前缀、BF状态、31维特征、五个邻居、预测和最终选择均完全一致，独立检查全部通过。\n'+f"四组中有{sum(g['oracle_preference_reverses'] for g in s['risk']['groups'])}组在不同未来下出现严格相反的oracle最优分支。合法前缀选择器不能因尚未发生的后缀而改变决定。风险集相对训练SBS平均多{s['risk']['summary']['selector_minus_SBS_mean']:.3f}箱；相对事后分支oracle平均后悔{s['risk']['summary']['selector_oracle_regret']:.3f}箱，最坏后悔{max(max(g['selector_regrets']) for g in s['risk']['groups'])}箱；这12流单列，不混入主指标，也不用于重新训练。")
sec('成本和独立整数验算',f"768次运行全完成，主验算覆盖{c['integer_and_decision_steps']:,}次容量、体积守恒、箱数和first-argmax决策。独立checker不导入runner或selector，用math标量评分、残量桶与最小索引堆重放控制决策，另写特征与kNN预测核对；120个部署选择及4组未来不变性均通过。\n数值执行总CPU {r['total_cpu_seconds']:.3f}秒、墙钟{r['total_wall_seconds']:.3f}秒；独立验算CPU {c['cpu_seconds']:.3f}秒；固定分析及制图CPU {a['cpu_seconds']:.3f}秒。外层记录三个成功阶段完整进程总CPU {totalcpu:.3f}秒，含启动和落盘；另为已停止的首次验证器计入240秒保守上界。标签、前缀、选择推断、序列化、训练和独立验算均有分项；不是将CPU隐匿到额外免费的oracle。主质量指标仍为箱数，CPU未被硬凑成箱数效用。\n单线程、新模型/API调用0、额外收费0。768个策略运行没有失败或重试。首次验证器因重复解压导致低效而停止26条后重启缓存版，全部原始日志保留；首次验证器CPU精确计量缺失，以其240秒硬限保守计入预算；执行前修正独立checker的浮点加法括号，并增加580秒软CPU异常保存；600秒硬限不变，v0/v1和修订证据保留，选择器和指标未变。跨平台浮点位级等价并未得到证明。\n历史整体900 CPU秒预算合规仍未知：首次验证器中断丢失精确计量，旧准备/静态检查也未完整计时。所留100秒只是计划预留，不能当作已证历史上界或实测成本。后续以396.212秒前瞻额度控制和计量；该会计限制与已通过的轨迹和结果正确性分开。")
sec('结果应如何使用','可以报告：在预冻的六尺寸族和三到达顺序的等权合成混合中，一个只见500项前缀的固定选择器取得上述可复查结果，含负例、同状态oracle后悔、已测成本及未完整计时项。\n不能报告：原132研究已通过新的确认；实现AutoFolio/RAISE全方法；完整流oracle可部署；任意未来、未知新分布或农业田间都享有同样收益。训练/确认同分布族、样本小且顺序结构固定，预测器可能依赖这种稳定性。只切换一次、只选两个固定控制；没有证明这是最强选择模型。\n本轮没有读取原132 heldout或P1结果，也没有改旧包、原提示、父代选择或原研究冻结。新数据与反馈不得接入原132第二代。')
sec('证据和重跑入口','PROTOCOL.json固定分布、特征、种子块、指标及预算；FREEZE.json固定执行前代码和输入。MODEL_FREEZE.json证明模型在确认运行前写入。inputs保存264条完整输入及谱系；raw保存768条完整分配/残量/可行数轨迹和成本；analysis保存逐流表、全部单元表、bootstrap重抽索引与分项CPU。README给出从空目录重跑的命令。COST_LEDGER.json补充完整进程、独立审阅和报告制作的实测成本；预冻结规划和少量只读命令没有完整计时，单独披露而不伪造为零。\nP069原文PDF与M04原件SHA256一致，关键机制参见PDF页6–7、10–14。https://ada.liacs.nl/papers/LinEtAl15.pdf\n固定EoH来源：Liu等，ICML 2024，Figure 6与§4.1。https://proceedings.mlr.press/v235/liu24bs.html')
md=['# BPP前缀安全算法选择实验','2026年10月6日']
for title,text in sections:md+=['## '+title,text]
md+=['## 全部确认单元','|尺寸族|顺序|全BF|BF后EoH|全EoH|选择器|选择器减SBS|选择器后悔|','|---|---|---:|---:|---:|---:|---:|---:|']
for row in rows:md.append('|'+ '|'.join([labels[row['family']],orders[row['order']]]+[f"{float(row[k]):.3f}" for k in ['BF','BF_then_EoH','published_EoH','SELECTOR','selector_minus_SBS','selector_regret']])+'|')
(H/'REPORT_ZH.md').write_text('\n\n'.join(md)+'\n')
doc=Document();page=doc.sections[0];page.page_height=Inches(11.7);page.page_width=Inches(8.3);page.top_margin=Inches(.65);page.bottom_margin=Inches(.65);page.left_margin=Inches(.7);page.right_margin=Inches(.7)
for name in ['Normal','Title','Heading 1','Heading 2']:
    st=doc.styles[name];st.font.name='Noto Sans CJK SC';st.font.color.rgb=RGBColor(0,0,0);st._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Noto Sans CJK SC')
doc.styles['Normal'].font.size=Pt(10.5);doc.styles['Normal'].paragraph_format.space_after=Pt(7);doc.styles['Normal'].paragraph_format.line_spacing=1.1;doc.styles['Title'].font.size=Pt(21);doc.styles['Heading 1'].font.size=Pt(14)
for st in doc.styles:
    for el in list(st._element.iter(qn('w:pBdr'))):el.getparent().remove(el)
def paragraphs(text):
    for p in text.split('\n'):doc.add_paragraph(p)
def table(headers,records,widths=None,font=9):
    t=doc.add_table(rows=1,cols=len(headers));t.style='Table Grid';t.autofit=False
    for cell,val in zip(t.rows[0].cells,headers):cell.text=str(val)
    for vals in records:
        for cell,val in zip(t.add_row().cells,vals):cell.text=str(val)
    for row in t.rows:
        row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
        for j,cell in enumerate(row.cells):
            if widths:cell.width=Inches(widths[j])
            for para in cell.paragraphs:
                para.paragraph_format.space_after=Pt(2);para.paragraph_format.space_before=Pt(2)
                for run in para.runs:run.font.size=Pt(font)
    t.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
    for cell in t.rows[0].cells:
        sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'E8EEF3');cell._tc.get_or_add_tcPr().append(sh)
    return t
# P1 conclusion and outcomes.
doc.add_paragraph('BPP前缀安全算法选择实验','Title');doc.add_paragraph('5000项完整执行  固定500项前缀  2026年10月6日')
doc.add_heading(sections[0][0],1);paragraphs(sections[0][1]);doc.add_picture(str(H/'figures/confirmation_means.png'),width=Inches(6.8));doc.add_paragraph('图1 均值来自108条确认流。oracle只是事后分支参照，不是额外可执行的选择器。')
# P2 methods.
doc.add_page_break()
for i in [1,2,3]:doc.add_heading(sections[i][0],1);paragraphs(sections[i][1])
# P3 all18 cells and heatmap.
doc.add_page_break();doc.add_heading(sections[4][0],1);paragraphs(sections[4][1])
table(['族','顺序','全BF','BF→EoH','全EoH','选择器','差SBS','后悔'],[[labels[x['family']],orders[x['order']]]+[f'{float(x[k]):.2f}' for k in ['BF','BF_then_EoH','published_EoH','SELECTOR','selector_minus_SBS','selector_regret']] for x in rows],[.85,.45,.9,.95,.9,.9,.8,.8],font=8)
doc.add_paragraph('差SBS为主指标的单元贡献，负值更好；后悔均相对同BF前缀分支oracle。完整逐流结果随包保留。')
doc.add_paragraph('对应热图另见figures/family_order_deltas.png，原始18单元值见analysis/family_order_summary.csv。')
# P4 prefix risk.
doc.add_page_break();doc.add_heading(sections[5][0],1);paragraphs(sections[5][1]);doc.add_picture(str(H/'figures/same_prefix_risk.png'),width=Inches(6.8))
table(['前缀族','固定选择','三种未来的后悔','oracle反转'],[[labels[g['group'].replace('risk_prefix_','')],pnames[g['choice']],'/'.join(map(str,g['selector_regrets'])),'是' if g['oracle_preference_reverses'] else '否'] for g in s['risk']['groups']],[1.05,2.1,2.4,1.1],9)
doc.add_paragraph('后悔顺序依次为随机、升序、降序后缀。即使后缀多重集相同，仅顺序变化也可能改变应选分支。')
# P5 cost and bounds.
doc.add_page_break();doc.add_heading(sections[6][0],1);paragraphs(sections[6][1])
table(['集合','方法','运行数','循环CPU秒','选择CPU秒'],[[{'train':'训练','confirmation':'确认','risk':'风险'}[x['split']],pnames[x['policy']],x['runs'],f"{float(x['cpu_loop_seconds']):.3f}",f"{float(x['cpu_selector_seconds']):.6f}"] for x in cpurows],[.6,2.6,.6,1.2,1.4],8.5)
doc.add_paragraph(f"模型拟合CPU {mf['fit_cpu_seconds']:.6f}秒。分项循环计时不含其后的压缩落盘，RUN_COMPLETE总量包含这些开销。用于审计的基线额外特征提取也单列在cpu_accounting.csv。")
# P6 reproducibility, explicit independent review status separate artifact.
doc.add_page_break();doc.add_heading(sections[7][0],1);paragraphs(sections[7][1]);doc.add_heading(sections[8][0],1);paragraphs(sections[8][1]);doc.add_heading('关键哈希',1)
for path in ['FREEZE.json','MODEL.json','PROTOCOL.json']:
    doc.add_paragraph(path+'\n'+hashlib.sha256((H/path).read_bytes()).hexdigest())
doc.add_paragraph('独立预审与结果复核文件一并归档。文档中未以oracle标签命中率代替实际箱数，也未将无差异解读为普适等效。')
doc.save(H/'BPP前缀安全算法选择实验.docx');print(json.dumps({'report_created':True,'experiment_cpu_seconds':totalcpu,'primary_delta':delta},ensure_ascii=False))
